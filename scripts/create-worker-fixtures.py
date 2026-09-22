from datetime import date
from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "reports" / "worker-tests"
OUT.mkdir(parents=True, exist_ok=True)
image = ROOT / "fixtures" / "ui-test" / "screenshots" / "desktop.png"


def create_docx():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)

    styles = doc.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(10)
    for name in ("Title", "Heading 1", "Heading 2"):
        styles[name].font.name = "Aptos Display"
        styles[name].font.color.rgb = RGBColor(0, 0, 0)
    styles["Title"].font.size = Pt(24)

    doc.add_heading("Setup Worker Smoke Test", 0)
    doc.add_paragraph(
        "Synthetic report proving document generation, screenshot inclusion, and structured output. "
        "No account, credential, or personal data is used."
    )
    doc.add_heading("Observed result", 1)
    doc.add_paragraph(
        "The disposable UI fixture rendered locally. Browser automation reached the page, read its controls, "
        "saved synthetic input, and confirmed the visible saved state."
    )
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    header = table.rows[0].cells
    header[0].text = "Check"
    header[1].text = "Result"
    header[2].text = "Evidence"
    for cell in header:
        for run in cell.paragraphs[0].runs:
            run.font.bold = True
        cell.vertical_alignment = 1
    for row in [
        ("Browser navigation", "PASS", "Local fixture loaded"),
        ("Information extraction", "PASS", "Heading, status, field, buttons"),
        ("Structured result", "PASS", "Saved synthetic name"),
        ("Screenshot", "PASS", "Actual fixture state"),
    ]:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            cells[idx].text = value
            cells[idx].vertical_alignment = 1
    doc.add_heading("Captured UI state", 1)
    if image.exists():
        doc.add_picture(str(image), width=Inches(6.2))
        doc.paragraphs[-1].alignment = 1
        doc.add_paragraph("Figure 1. Disposable fixture screenshot supplied by the local test run.").alignment = 1
    doc.add_heading("Source note", 1)
    doc.add_paragraph("Source: fixtures/ui-test/index.html and the local browser smoke test. Date: 2026-09-20.")
    path = OUT / "worker-smoke-test.docx"
    doc.save(path)
    return path


def create_xlsx():
    wb = Workbook()
    ws = wb.active
    ws.title = "Synthetic Research"
    headers = ["Company", "Role", "Fit score", "Status", "Source URL", "Date found"]
    ws.append(headers)
    rows = [
        ["Northwind Labs", "Product Engineer", 0.92, "Shortlist", "https://example.com/northwind", date(2026, 9, 20)],
        ["Fabrikam Systems", "QA Engineer", 0.78, "Review", "https://example.com/fabrikam", date(2026, 9, 20)],
        ["Contoso Research", "Automation Engineer", 0.86, "Shortlist", "https://example.com/contoso", date(2026, 9, 20)],
    ]
    for row in rows:
        ws.append(row)
    ws.append([])
    ws.append(["Average fit score", "=AVERAGE(C2:C4)"])
    ws.append(["Shortlist count", '=COUNTIF(D2:D4,"Shortlist")'])
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center")
    for row in ws.iter_rows(min_row=2, max_row=4, min_col=3, max_col=3):
        row[0].number_format = "0%"
    for row in ws.iter_rows(min_row=2, max_row=4, min_col=6, max_col=6):
        row[0].number_format = "yyyy-mm-dd"
    ws["B6"].number_format = "0%"
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = "A1:F4"
    widths = {"A": 22, "B": 24, "C": 12, "D": 14, "E": 36, "F": 14}
    for column, width in widths.items():
        ws.column_dimensions[column].width = width
    path = OUT / "synthetic-research.xlsx"
    wb.save(path)
    # Reopen to prove exported workbook is readable and formulas are present.
    check = load_workbook(path, data_only=False)
    assert check["Synthetic Research"]["B6"].value == "=AVERAGE(C2:C4)"
    assert check["Synthetic Research"]["B7"].value == '=COUNTIF(D2:D4,"Shortlist")'
    return path


if __name__ == "__main__":
    print(create_docx())
    print(create_xlsx())
